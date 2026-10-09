"""Resume document parsing (PDF / DOCX) with friendly error messages."""
import io
import logging

logger = logging.getLogger("careermirror.parser")


class ResumeParseError(Exception):
    """Raised when a document cannot be parsed. Message is user-facing."""


def parse_pdf(content: bytes) -> str:
    try:
        import pdfplumber
    except ImportError:
        raise ResumeParseError("PDF support is not installed. Run: pip install pdfplumber")
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = []
            for page in pdf.pages:
                text = page.extract_text() or ""
                pages.append(text)
            full = "\n".join(pages)
    except Exception as exc:
        raise ResumeParseError(
            f"Resume could not be parsed. The PDF may be corrupted, password-protected "
            f"or image-only. Details: {exc}"
        ) from exc
    if not full.strip():
        raise ResumeParseError(
            "No text could be extracted from the PDF. It may be a scanned/image-only "
            "document - please upload a text-based PDF or a DOCX file."
        )
    return full


def parse_docx(content: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise ResumeParseError("DOCX support is not installed. Run: pip install python-docx")
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ResumeParseError(
            f"Resume could not be parsed. The DOCX file may be corrupted. Details: {exc}"
        ) from exc

    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    full = "\n".join(parts)
    if not full.strip():
        raise ResumeParseError(
            "No text could be extracted from the DOCX file. The document may be empty "
            "or contain only images."
        )
    return full


def parse_resume(content: bytes, extension: str) -> str:
    """Parse a resume file into plain text."""
    if extension == ".pdf":
        return parse_pdf(content)
    if extension == ".docx":
        return parse_docx(content)
    raise ResumeParseError("Resume must be a PDF or DOCX file.")
