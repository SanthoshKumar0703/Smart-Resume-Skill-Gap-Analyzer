"""Generate sample resumes (PDF + DOCX) and a sample job description
so the app can be tested end-to-end without personal data."""
import io
import os
import sys

RESUME_TEXT = """AARAV SHARMA
Bengaluru, India | aarav.sharma@example.com | +91 98765 43210
linkedin.com/in/aaravsharma | github.com/aaravsharma

PROFESSIONAL SUMMARY
Full-stack developer with 4+ years of experience building web applications with
Python and JavaScript. Strong background in React frontends, REST API design and
MongoDB data modeling. Known for strong communication and teamwork skills,
mentoring juniors and speaking at internal tech events.

TECHNICAL SKILLS
Python, JavaScript, TypeScript, React, Redux, HTML5, CSS3, Tailwind CSS, Node.js,
Express, Flask, REST APIs, JWT authentication, MongoDB, PostgreSQL, Redis,
Git, GitHub Actions, CI/CD, Jest, Linux, Agile, Scrum

WORK EXPERIENCE
Software Engineer | TechNova Solutions, Bengaluru
Jan 2022 - Present
- Built and maintained React + Redux dashboards serving 200k monthly users
- Designed REST APIs in Python (Flask) with JWT auth and role-based access
- Modeled MongoDB schemas and aggregation pipelines for analytics features
- Set up GitHub Actions CI pipelines with automated Jest test suites
- Introduced Redis caching, cutting API latency by 40%

Junior Developer | WebWorks Studio, Pune
Jun 2019 - Dec 2021
- Developed responsive UIs with React, Tailwind CSS and TypeScript
- Implemented CRUD backends with Node.js, Express and PostgreSQL
- Wrote unit tests with Jest; participated in daily scrum ceremonies
- Refactored legacy jQuery code to React, improving page speed by 35%

EDUCATION
B.Tech in Computer Science and Engineering
Visvesvaraya National Institute of Technology, 2015 - 2019

PROJECTS
- SkillBridge: A MERN stack platform connecting mentors and mentees with chat and
  video scheduling; deployed on Heroku with Redis for presence tracking
- Resume Lens: Python tool that extracts structured data from resumes using NLP,
  with a React + TypeScript UI and MongoDB backend
- DevStats: GitHub analytics dashboard aggregating commit data via REST APIs

CERTIFICATIONS
Meta Front-End Developer Professional Certificate (Coursera)

ACHIEVEMENTS
- Speaker at internal tech talks on React performance patterns
"""

JD_TEXT = """Full Stack Developer (React + Python)

About the role
We are looking for a Full Stack Developer to join our product engineering team.
You will build scalable web applications, design APIs, and ship features used by
millions of users. You will work closely with product and design teams in an
agile environment.

Responsibilities
- Build responsive, accessible frontends using React and modern JavaScript
- Design and implement REST APIs with Python (FastAPI preferred)
- Model and optimize data with MongoDB and PostgreSQL
- Containerize services with Docker and orchestrate with Kubernetes
- Deploy and manage infrastructure on AWS (EC2, S3, ECS)
- Write automated tests and maintain CI/CD pipelines
- Participate in code reviews, scrum ceremonies and technical design sessions

Requirements
- 3+ years of professional software development experience
- Strong proficiency in Python, React, JavaScript and TypeScript
- Hands-on experience with MongoDB and SQL databases
- Experience building REST APIs with FastAPI or similar Python frameworks
- Working knowledge of Docker, AWS and Kubernetes
- Familiarity with Redis, Git and CI/CD practices
- Bachelor's degree in Computer Science or a related field
- Excellent communication and teamwork skills

Nice to have
- AWS Certified Developer Associate or equivalent
- Experience with microservices architecture and system design
- Experience with GraphQL
- Contributions to open-source projects
"""


def build_docx(path):
    import docx
    doc = docx.Document()
    doc.add_heading("AARAV SHARMA", 0)
    for line in RESUME_TEXT.split("\n")[1:]:
        line = line.strip()
        if not line:
            continue
        if line.isupper() or (line.endswith(":") and len(line) < 45):
            doc.add_heading(line.title(), level=1)
        elif line.startswith(("- ", "• ")):
            doc.add_paragraph(line, style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.save(path)


def build_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.lib.enums import TA_CENTER

    styles = {
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=18,
                                alignment=TA_CENTER, spaceAfter=2),
        "sub": ParagraphStyle("sub", fontName="Helvetica", fontSize=8.5,
                              alignment=TA_CENTER, textColor="#555555", spaceAfter=8),
        "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11,
                            spaceBefore=10, spaceAfter=3, textColor="#0e5a44"),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=9.2,
                               leading=13, spaceAfter=2),
    }
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm)
    story = [Paragraph("AARAV SHARMA", styles["title"]),
             Paragraph("Bengaluru, India | aarav.sharma@example.com | +91 98765 43210<br/>linkedin.com/in/aaravsharma | github.com/aaravsharma", styles["sub"])]
    lines = RESUME_TEXT.split("\n")[1:]
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.isupper() or (line.endswith(":") and len(line) < 45):
            story.append(Paragraph(line.title(), styles["h"]))
        elif line.startswith("- "):
            story.append(Paragraph("- " + line[2:], styles["body"]))
        else:
            story.append(Paragraph(line, styles["body"]))
    doc.build(story)


def main():
    out_dir = os.path.join(os.path.dirname(__file__), "..", "..", "samples")
    os.makedirs(out_dir, exist_ok=True)
    build_docx(os.path.join(out_dir, "sample_resume.docx"))
    build_pdf(os.path.join(out_dir, "sample_resume.pdf"))
    with open(os.path.join(out_dir, "sample_job_description.txt"), "w", encoding="utf-8") as f:
        f.write(JD_TEXT)
    print("Samples written to", os.path.abspath(out_dir))


if __name__ == "__main__":
    main()
